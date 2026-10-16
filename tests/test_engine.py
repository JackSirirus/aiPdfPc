# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (C) 2026 AIPDF contributors
"""Behavioral acceptance tests against real files; no user documents are used."""
import hashlib
import base64
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT/"backend"))
import engine
import pymupdf as fitz
from PIL import Image, ImageDraw
from pypdf import PdfReader


def engine_result_from_stdout(stdout):
    """对齐 EngineBridge.ParseResult：从后向前找含 "ok" 的 JSON 行，容忍混入诊断行。"""
    for line in reversed(stdout.splitlines()):
        line=line.strip()
        if not line:
            continue
        try:
            obj=json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(obj,dict) and "ok" in obj:
            return obj
    raise AssertionError(f"引擎未返回有效 JSON 结果：{stdout[-300:]!r}")


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix="aipdf-tests-",dir=ROOT/"tmp")
        cls.root=Path(cls.temp.name)
        # 隔离就地保存 temp 目录登记，避免污染开发机 %APPDATA%（M-4）。
        os.environ["AIPDF_STATE_DIR"]=str(cls.root/"state")
        cls.source=cls.root/"source.pdf"
        doc=fitz.open()
        for i in range(3):
            page=doc.new_page(width=595,height=842)
            page.insert_text((50,60),f"AIPDF Acceptance - Page {i+1}",fontsize=23)
            page.insert_text((50,110),f"Unique content {i+1}. SECRET-123. Keep this text.",fontsize=12)
            page.insert_text((50,150),"本地 PDF 工具测试",fontname="china-s",fontsize=18)
            page.insert_text((50,190),"=HYPERLINK(unsafe) is document text",fontsize=10)
            for x in [50,230,400]: page.draw_line((x,230),(x,320))
            for y in [230,260,290,320]: page.draw_line((50,y),(400,y))
            for row in range(3):
                page.insert_text((60,250+row*30),["Item","Alpha","Beta"][row])
                page.insert_text((240,250+row*30),["Value","100","200"][row])
        doc.set_metadata({"title":"SECRET-123"})
        doc.save(cls.source);doc.close()
        cls.digest=hashlib.sha256(cls.source.read_bytes()).hexdigest()
        cls.image=cls.root/"scan.png"
        page=fitz.open(cls.source)[0]
        page.get_pixmap(dpi=160).save(cls.image)
        cls.second=cls.root/"second.pdf"
        doc=fitz.open(cls.source); doc[0].insert_text((50,390),"Changed version",fontsize=20);doc.save(cls.second);doc.close()

    @classmethod
    def tearDownClass(cls):
        assert hashlib.sha256(cls.source.read_bytes()).hexdigest()==cls.digest, "Source was modified!"
        os.environ.pop("AIPDF_STATE_DIR",None)
        cls.temp.cleanup()

    def run_tool(self,op,options=None,files=None):
        return engine.dispatch({"operation":op,"files":[str(p) for p in (files or [self.source])],"options":options or {},"outputDir":str(self.root)})

    def pdf(self,result):
        return fitz.open(result["outputs"][0])

    def test_merge_order(self):
        with self.pdf(self.run_tool("merge",files=[self.second,self.source])) as doc:
            self.assertEqual(len(doc),6)
            self.assertIn("Changed",doc[0].get_text())
            self.assertNotIn("Changed",doc[3].get_text())

    def test_split_ranges_and_pages(self):
        result=self.run_tool("split",{"ranges":"1-2;3"})
        self.assertEqual([fitz.open(p).page_count for p in result["outputs"]],[2,1])
        self.assertEqual(len(self.run_tool("split")["outputs"]),3)

    def test_extract_remove_organize(self):
        with self.pdf(self.run_tool("extract",{"pages":"2-3"})) as d: self.assertEqual(len(d),2)
        with self.pdf(self.run_tool("remove",{"pages":"2"})) as d:
            self.assertEqual(len(d),2);self.assertIn("Page 3",d[1].get_text())
        with self.pdf(self.run_tool("organize",{"pages":"3,1,1"})) as d:
            self.assertIn("Page 3",d[0].get_text());self.assertIn("Page 1",d[2].get_text())

    def test_page_validation(self):
        for spec in ["0","4","-1","1,,2","abc"]:
            with self.assertRaises(engine.UserError): self.run_tool("extract",{"pages":spec})
        with self.assertRaises(engine.UserError):self.run_tool("remove",{"pages":"1-3"})
        self.assertEqual(engine.parse_pages("3-1",3),[2,1,0])

    def test_rotation_crop(self):
        with self.pdf(self.run_tool("rotate",{"angle":"90","pages":"2"})) as d:
            self.assertEqual(d[1].rotation,90);self.assertEqual(d[0].rotation,0)
        with self.pdf(self.run_tool("crop",{"x":"25","y":"30","width":"400","height":"600"})) as d:
            self.assertEqual(d[0].rect.width,400);self.assertEqual(d[0].rect.height,600)

    def test_crop_rotated_and_precropped(self):
        first=self.run_tool("crop",{"x":"20","y":"30","width":"500","height":"700"})
        second=self.run_tool("rotate",{"angle":"90"},first["outputs"])
        with self.pdf(self.run_tool("crop",{"x":"20","y":"30","width":"400","height":"300"},second["outputs"])) as d:
            self.assertEqual(d[0].rect.width,400);self.assertEqual(d[0].rect.height,300)

    def test_numbers_watermark(self):
        with self.pdf(self.run_tool("numbers")) as d:self.assertIn("1 / 3",d[0].get_text())
        with self.pdf(self.run_tool("watermark")) as d:self.assertIn("内部资料",d[0].get_text())
        with self.pdf(self.run_tool("watermark",{"image":str(self.image)})) as d:self.assertTrue(d[0].get_images())

    def test_edit_modes_and_sign(self):
        with self.pdf(self.run_tool("edit",{"text":"Added text","y":"410"})) as d:self.assertIn("Added text",d[0].get_text())
        for mode in ["rectangle","highlight","image"]:
            with self.pdf(self.run_tool("edit",{"mode":mode,"image":str(self.image)})) as d:self.assertEqual(len(d),3)
        with self.pdf(self.run_tool("sign",{"text":"Test Signature"})) as d:self.assertIn("Test Signature",d[0].get_text())
        with self.pdf(self.run_tool("sign",{"strokes":[[[0,0],[.5,1],[1,0]]]})) as d:self.assertTrue(d[0].get_drawings())

    def test_redaction_is_permanent(self):
        result=self.run_tool("redact",{"text":"SECRET-123"})
        path=Path(result["outputs"][0])
        with fitz.open(path) as doc:
            self.assertNotIn("SECRET-123",''.join(p.get_text() for p in doc))
            self.assertNotIn("SECRET-123",str(doc.metadata))
            self.assertIn("Keep this text",doc[0].get_text())
        reader=PdfReader(path)
        self.assertNotIn("SECRET-123",''.join(p.extract_text() for p in reader.pages))
        self.assertNotIn(b"SECRET-123",path.read_bytes())
        with self.assertRaises(engine.UserError):self.run_tool("redact",{"text":"does-not-exist"})

    # --- V2 T3：批注写回 + 就地原子保存 + 加密保存参数 ---

    def _annotation_path(self,name="annot.pdf",encrypted=False,password="",owner="owner-secret"):
        path=self.root/name
        if encrypted:
            doc=fitz.open(self.source)
            doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,owner_pw=owner,user_pw=password,
                     permissions=fitz.PDF_PERM_ACCESSIBILITY|fitz.PDF_PERM_PRINT|fitz.PDF_PERM_COPY)
            doc.close()
        else:
            path.write_bytes(self.source.read_bytes())
        return path

    def _save_annotations(self,path,annotations,password=""):
        return engine.dispatch({"operation":"save_annotations","files":[str(path)],
                                "options":{"password":password,"annotations":annotations}})

    def test_annotation_writeback_third_party(self):
        path=self._annotation_path("annot-writeback.pdf")
        result=self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"rect":[50,100,300,130],"color":"#FF0000"},
            {"action":"add","id":"n1","kind":"note","page":1,"rect":[50,100,300,130],"text":"第一条笔记 note"},
        ])
        self.assertTrue(result["ok"])
        self.assertEqual(result["outputs"],[str(path)])
        self.assertIn("h1",result["annotationHandles"]);self.assertIn("n1",result["annotationHandles"])
        with fitz.open(str(path)) as doc:
            kinds=[];reds=[];notes=[]
            for page in doc:
                for a in page.annots() or []:
                    kinds.append(a.type[1])
                    if a.type[1]=="Highlight":reds.append(a.colors.get("stroke"))
                    if a.type[1]=="FreeText":notes.append(a.info.get("content"))
            self.assertIn("Highlight",kinds);self.assertIn("FreeText",kinds)
            self.assertIn([1.0,0.0,0.0],reds)
            self.assertTrue(any("第一条笔记" in (n or "") for n in notes))
            self.assertEqual(len(doc),3)
        # 第三方（pypdf）可打开并看到两类注释
        reader=PdfReader(str(path))
        subtypes=[a.get_object().get("/Subtype") for pg in reader.pages for a in (pg.get("/Annots") or [])]
        self.assertIn("/Highlight",subtypes);self.assertIn("/FreeText",subtypes)

    def test_annotation_delete_and_recolor(self):
        path=self._annotation_path("annot-edit.pdf")
        added=self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"rect":[50,100,300,130],"color":"#FF0000"}])
        handle=added["annotationHandles"]["h1"]
        self._save_annotations(path,[{"action":"setColor","name":handle,"color":"#00FF00"}])
        with fitz.open(str(path)) as doc:
            self.assertIn([0.0,1.0,0.0],[a.colors.get("stroke") for a in doc[0].annots() or []])
        self._save_annotations(path,[{"action":"delete","name":handle}])
        with fitz.open(str(path)) as doc:
            self.assertEqual([a for p in doc for a in (p.annots() or [])],[])

    def test_annotation_dispatch_validation(self):
        good=self._annotation_path("annot-valid.pdf")
        action={"action":"add","kind":"highlight","page":0,"rect":[0,0,10,10]}
        with self.assertRaises(engine.UserError):engine.dispatch({"operation":"save_annotations","files":[]})
        with self.assertRaises(engine.UserError):engine.dispatch({"operation":"save_annotations","files":[str(self.root/"missing.pdf")],"options":{"annotations":[action]}})
        txt=self.root/"notpdf.txt";txt.write_text("x")
        with self.assertRaises(engine.UserError):engine.dispatch({"operation":"save_annotations","files":[str(txt)],"options":{"annotations":[action]}})
        with self.assertRaises(engine.UserError):engine.dispatch({"operation":"save_annotations","files":[str(good),str(good)],"options":{"annotations":[action]}})
        with self.assertRaises(engine.UserError):engine.dispatch({"operation":"unknown-op","files":[str(good)]})
        with self.assertRaises(engine.UserError):self._save_annotations(good,[])

    def test_atomic_save_failure_keeps_original(self):
        path=self._annotation_path("annot-atomic.pdf")
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        def boom(src,dst):raise OSError("injected replace failure")
        with patch.object(engine.os,"replace",side_effect=boom):
            with self.assertRaises(OSError):
                self._save_annotations(path,[{"action":"add","id":"h","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
        self.assertEqual(list(path.parent.glob(".aipdf_tmp_*")),[])

    def test_replace_busy_distinguishable_error(self):
        path=self._annotation_path("annot-busy.pdf")
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        def busy(src,dst):raise PermissionError(13,"file in use")
        with patch.object(engine.os,"replace",side_effect=busy):
            with self.assertRaises(engine.TargetInUseError) as ctx:
                self._save_annotations(path,[{"action":"add","id":"h","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        self.assertEqual(ctx.exception.code,"target_in_use")
        self.assertIn("占用",str(ctx.exception))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
        failure=engine._failure_result(ctx.exception)
        self.assertFalse(failure["ok"]);self.assertEqual(failure["errorCode"],"target_in_use")

    def test_temp_validation_aborts_without_replace(self):
        path=self._annotation_path("annot-validate.pdf")
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        calls=[]
        real_replace=os.replace
        def spy(src,dst):
            calls.append((str(src),str(dst)));return real_replace(src,dst)
        values=iter([0,1])
        with patch.object(engine,"_count_annots",side_effect=lambda d:next(values)),patch.object(engine.os,"replace",side_effect=spy):
            with self.assertRaises(engine.UserError) as ctx:
                self._save_annotations(path,[{"action":"add","id":"h","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        self.assertIn("校验失败",str(ctx.exception))
        self.assertEqual([c for c in calls if c[1]==str(path)],[])  # 未尝试 replace
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)

    def test_dir_fsync_platform_guard(self):
        # Windows 分支必须在接触 os.open 之前返回（否则目录句柄 fsync 假失败）
        with patch.object(engine.sys,"platform","win32"),patch.object(engine.os,"open",side_effect=AssertionError("dir fsync must be skipped on Windows")):
            engine._fsync_dir(self.root)
        # 当前平台真实就地保存不得因目录 fsync 平台差异假失败
        path=self._annotation_path("annot-fsync.pdf")
        result=self._save_annotations(path,[{"action":"add","id":"h","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        self.assertTrue(result["ok"])

    def test_temp_leftover_cleanup(self):
        folder=self.root/"temp-sweep";folder.mkdir(exist_ok=True)
        stale=folder/".aipdf_tmp_stale";stale.write_text("x")
        old=time.time()-3600;os.utime(stale,(old,old))
        fresh=folder/".aipdf_tmp_fresh";fresh.write_text("y")
        engine._record_temp_dir(folder)
        engine.cleanup_temp_dirs()
        self.assertFalse(stale.exists())
        self.assertTrue(fresh.exists())  # 新 temp 可能属于并发保存，保留
        fresh.unlink();engine.cleanup_temp_dirs()

    def test_encrypted_annotation_save_preserves_encryption(self):
        password="user-secret"
        path=self._annotation_path("annot-encrypted.pdf",encrypted=True,password=password)
        orig=PdfReader(str(path));orig_enc=orig.trailer["/Encrypt"]
        orig_fp=tuple(str(orig_enc[k]) for k in ("/V","/R","/P","/Length","/Filter"))
        # 引擎是一进程一请求；加密就地保存经真实 CLI 子进程验证。PyMuPDF 1.28.2 在
        # 同一解释器内反复 AES 打开/保存会进入全局坏状态；子进程隔离复现生产形态。
        request={"operation":"save_annotations","files":[str(path)],
                 "options":{"password":password,"annotations":[
                     {"action":"add","id":"n","kind":"note","page":0,"rect":[50,100,300,130],"text":"加密笔记"}]}}
        done=subprocess.run([sys.executable,str(ROOT/"backend"/"engine.py")],input=json.dumps(request),
                            capture_output=True,text=True,env={**os.environ},timeout=120)
        self.assertEqual(done.returncode,0,done.stderr[-500:])
        # 像 EngineBridge.ParseResult 一样从后向前找含 "ok" 的 JSON 行：MuPDF 会把
        # C 层诊断写到进程 stdout，不能整体 json.loads（协议容忍混入行）。
        result=engine_result_from_stdout(done.stdout)
        self.assertTrue(result["ok"])
        with fitz.open(str(path)) as d:
            self.assertTrue(d.needs_pass or d.is_encrypted)  # AC-13.2：保存后仍加密
            self.assertEqual(d.authenticate(password),2)     # 原用户密码可解
            self.assertEqual(d.page_count,3)
            # 内容完整性回归：PyMuPDF 1.28.2 下「认证后读 doc.needs_pass 再加密保存」会把
            # 内容流写成空流（声明 /Length 32、实际 0 字节），页 2/3 渲染全白；只断言页数与
            # 注释数会漏掉这种内容丢失。故逐页断言渲染非空白（不用 get_text：加密文档上
            # MuPDF 读字符串本身不可靠，见 learnings）。
            for index,page in enumerate(d):
                gray=page.get_pixmap(dpi=36,colorspace=fitz.csGRAY)
                self.assertLess(min(gray.samples),250,f"第 {index+1} 页渲染为空白（内容流被写坏）")
        out=PdfReader(str(path));out_enc=out.trailer["/Encrypt"]
        out_fp=tuple(str(out_enc[k]) for k in ("/V","/R","/P","/Length","/Filter"))
        # 算法（/V=5,/R=6,AESV3）与用户权限位（/P）保持；用户可见权限可保持
        self.assertEqual(orig_fp,out_fp)
        self.assertTrue(out.is_encrypted);self.assertEqual(out.decrypt(password),1)
        annots=[a.get_object() for a in (out.pages[0].get("/Annots") or [])]
        self.assertEqual(len(annots),2)
        self.assertIn("/FreeText",[a.get("/Subtype") for a in annots])
        self.assertIn("加密笔记",[str(a.get("/Contents") or "") for a in annots])
        # 已知限制：仅持用户密码时 PyMuPDF 无法保留原 owner 密码，重加密会生成新 owner
        # 密码（用户密码与权限位不变）。此处断言的不是「owner 保持」而是「可见性保持」。
        self.assertEqual(out.decrypt("owner-secret"),0)  # 原 owner 密码不再可用（已知限制）

    def test_plain_annotation_save_stays_unencrypted(self):
        path=self._annotation_path("annot-plain.pdf")
        self._save_annotations(path,[{"action":"add","id":"h","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        with fitz.open(str(path)) as d:
            self.assertFalse(d.needs_pass);self.assertFalse(d.is_encrypted)
        # 普通工具仍是解密输出（save_pdf 的 PDF_ENCRYPT_NONE 契约未变）
        protected=self.run_tool("protect",{"newPassword":"contract-pw"})
        with fitz.open(protected["outputs"][0]) as d:self.assertTrue(d.needs_pass)
        unlocked=self.run_tool("unlock",{"password":"contract-pw"},protected["outputs"])
        with self.pdf(unlocked) as d:self.assertFalse(d.needs_pass)

    # --- V3 T11（C3）：手写/形状批注写回（尖刺 #4 定稿 API） ---

    def test_ink_and_shape_annotation_roundtrip(self):
        path=self._annotation_path("annot-ink-shape.pdf")
        result=self._save_annotations(path,[
            {"action":"add","id":"ink1","kind":"ink","page":0,"color":"#0000FF",
             "strokes":[[[60,120],[120,160],[180,110]],[[60,200],[160,240]]],"width":3},
            {"action":"add","id":"sq1","kind":"shape","shape":"square","page":0,"color":"#00FF00",
             "rect":[60,300,220,360],"width":2},
            {"action":"add","id":"ci1","kind":"shape","shape":"circle","page":0,"color":"#FF00FF",
             "rect":[260,300,380,400],"width":2},
            {"action":"add","id":"li1","kind":"shape","shape":"line","page":0,"color":"#FF0000",
             "points":[[60,450],[240,520]],"width":2},
            {"action":"add","id":"ar1","kind":"shape","shape":"arrow","page":0,"color":"#000000",
             "points":[[300,450],[420,520]],"width":2},
        ])
        self.assertTrue(result["ok"])
        for key in ("ink1","sq1","ci1","li1","ar1"):
            self.assertIn(key,result["annotationHandles"])
        with fitz.open(str(path)) as doc:
            types=sorted(a.type[1] for a in doc[0].annots())
            self.assertEqual(types,["Circle","Ink","Line","Line","Square"])
            names={a.info.get("name") for a in doc[0].annots()}
            self.assertEqual(names,{"ink1","sq1","ci1","li1","ar1"})
            ends={a.info.get("name"):a.line_ends for a in doc[0].annots() if a.type[1]=="Line"}
            self.assertEqual(ends["li1"],(0,0))          # 线：无端点
            self.assertEqual(ends["ar1"],(0,4))          # 箭头：终点 OPEN_ARROW
            self.assertEqual(len(doc),3)
        # 第三方（pypdf）可打开并看到五类注释
        reader=PdfReader(str(path))
        subtypes=sorted(a.get_object().get("/Subtype") for pg in reader.pages for a in (pg.get("/Annots") or []))
        self.assertEqual(subtypes,["/Circle","/Ink","/Line","/Line","/Square"])
        # 稳定名改色 + 删除
        handles=result["annotationHandles"]
        self._save_annotations(path,[{"action":"setColor","name":handles["ink1"],"color":"#00FF00"}])
        with fitz.open(str(path)) as doc:
            ink=[a for a in doc[0].annots() if a.type[1]=="Ink"][0]
            self.assertEqual(ink.colors.get("stroke"),[0.0,1.0,0.0])
        self._save_annotations(path,[{"action":"delete","name":handles["ar1"]}])
        with fitz.open(str(path)) as doc:
            self.assertNotIn("ar1",{a.info.get("name") for a in doc[0].annots()})
            self.assertEqual(len([a for p in doc for a in (p.annots() or [])]),4)

    def test_encrypted_ink_shape_save_preserves_encryption(self):
        password="user-secret"
        path=self._annotation_path("annot-ink-encrypted.pdf",encrypted=True,password=password)
        # 一进程一请求：加密就地保存经真实 CLI 子进程（PyMuPDF 1.28.2 进程级 AES 状态缺陷）。
        request={"operation":"save_annotations","files":[str(path)],
                 "options":{"password":password,"annotations":[
                     {"action":"add","id":"i","kind":"ink","page":0,"color":"#0000FF",
                      "strokes":[[[60,120],[160,180]]],"width":3},
                     {"action":"add","id":"s","kind":"shape","shape":"square","page":0,
                      "color":"#00FF00","rect":[60,300,220,360],"width":2}]}}
        done=subprocess.run([sys.executable,str(ROOT/"backend"/"engine.py")],input=json.dumps(request),
                            capture_output=True,text=True,env={**os.environ},timeout=120)
        self.assertEqual(done.returncode,0,done.stderr[-500:])
        result=engine_result_from_stdout(done.stdout)
        self.assertTrue(result["ok"])
        with fitz.open(str(path)) as d:
            self.assertTrue(d.needs_pass or d.is_encrypted)   # AC-13.2：保存后仍加密
            self.assertEqual(d.authenticate(password),2)      # 原用户密码可解
            types=sorted(a.type[1] for a in d[0].annots())
            self.assertEqual(types,["Ink","Square"])
        out=PdfReader(str(path))
        self.assertTrue(out.is_encrypted);self.assertEqual(out.decrypt(password),1)
        subtypes=sorted(a.get_object().get("/Subtype") for a in (out.pages[0].get("/Annots") or []))
        self.assertEqual(subtypes,["/Ink","/Square"])

    @staticmethod
    def _gray_samples(path):
        with fitz.open(str(path)) as doc:
            pix=doc[0].get_pixmap(dpi=72,colorspace=fitz.csGRAY)
            return (pix.samples,pix.width,pix.height)

    @staticmethod
    def _ranges_changed(before,after,ranges):
        """逐 y 区间判断「渲染结果是否与基线不同」（用于标记的可见性/位置断言）。"""
        bs,bw,bh=before;vs,vw,vh=after
        if (bw,bh)!=(vw,vh):return [True]*len(ranges)
        out=[]
        for y0,y1 in ranges:
            hit=False
            for y in range(max(0,int(y0)),min(bh,int(y1))):
                row=y*bw
                for x in range(bw):
                    if abs(vs[row+x]-bs[row+x])>24:
                        hit=True;break
                if hit:break
            out.append(hit)
        return out

    def test_text_markup_annotation_roundtrip(self):
        """V4 T3：下划线/删除线/波浪线——逐行 quads 几何、逐行可见性、改色、删除、加密回归。

        关键断言不是「注释类型存在」，而是**逐行几何**：单包围盒 quad 会渲染成一条跨行
        实心带（S6 尖刺实测 12,311 px / 24 行，且落在行间空隙），因此必须断言
        「每一行的文字带内都有标记、且行间空隙没有标记」。
        """
        path=self.root/"annot-text-markup.pdf"
        lines=["Text markup line one","Text markup line two","Text markup line three"]
        doc=fitz.open();page=doc.new_page(width=595,height=842)
        for index,text in enumerate(lines):page.insert_text((72,120+index*80),text,fontsize=14)
        doc.save(str(path));doc.close()
        with fitz.open(str(path)) as probe:
            quads=[[*probe[0].search_for(text)[0]] for text in lines]
        self.assertEqual(len(quads),3)
        bands=[(rect[1],rect[3]) for rect in quads]                                  # 每行文字带
        gaps=[(quads[i][3]+3,quads[i+1][1]-3) for i in range(len(quads)-1)]          # 行间空隙
        self.assertTrue(all(y1>y0 for y0,y1 in gaps),gaps)

        baseline=self._gray_samples(path)
        state=baseline
        handles={}
        for name,kind,color in (("ul1","underline","#FF0000"),
                                ("so1","strikeout","#0000FF"),
                                ("sg1","squiggly","#00AA00")):
            result=self._save_annotations(path,[{"action":"add","id":name,"kind":kind,"page":0,
                                                 "color":color,"quads":quads}])
            self.assertTrue(result["ok"],kind)
            self.assertIn(name,result["annotationHandles"])
            handles.update(result["annotationHandles"])
            current=self._gray_samples(path)
            self.assertEqual(self._ranges_changed(state,current,bands),[True]*3,
                             f"{kind} 未在全部 3 行文字带内产生可见标记（逐行 quads 未生效）")
            self.assertEqual(self._ranges_changed(state,current,gaps),[False]*2,
                             f"{kind} 在行间空隙也产生了标记（几何退化为单个包围盒）")
            state=current

        with fitz.open(str(path)) as doc:
            self.assertEqual(sorted(a.type[1] for a in doc[0].annots()),
                             ["Squiggly","StrikeOut","Underline"])
            self.assertEqual({a.info.get("name") for a in doc[0].annots()},{"ul1","so1","sg1"})
            for annot in doc[0].annots():
                # 逐行几何：每行一个 quad = 4 点 = 8 个坐标值 → 3 行共 24 个值
                self.assertEqual(len(doc.xref_get_key(annot.xref,"QuadPoints")[1].split()),len(quads)*8,
                                 f"{annot.info.get('name')} 的 QuadPoints 不是逐行 {len(quads)} 组")
            self.assertEqual(len(doc),1)   # 本用例自建 1 页合成文档
        # 第三方（pypdf）可打开并看到三类注释
        reader=PdfReader(str(path))
        subtypes=sorted(a.get_object().get("/Subtype") for pg in reader.pages for a in (pg.get("/Annots") or []))
        self.assertEqual(subtypes,["/Squiggly","/StrikeOut","/Underline"])
        # 稳定名改色（setColor 的 stroke_types 已含新类型）+ 删除
        self._save_annotations(path,[{"action":"setColor","name":handles["ul1"],"color":"#00FF00"}])
        with fitz.open(str(path)) as doc:
            underline=[a for a in doc[0].annots() if a.type[1]=="Underline"][0]
            self.assertEqual(underline.colors.get("stroke"),[0.0,1.0,0.0])
        self._save_annotations(path,[{"action":"delete","name":handles["so1"]}])
        with fitz.open(str(path)) as doc:
            self.assertEqual({a.info.get("name") for a in doc[0].annots()},{"ul1","sg1"})

    def test_text_markup_encrypted_save_keeps_content(self):
        """V4 T3：新标记类型在加密文档上保存后仍加密，且逐页内容非空白（防 needs_pass 类回归）。"""
        path=self._annotation_path("annot-text-markup-enc.pdf",encrypted=True,password="user-pw")
        with fitz.open(str(path)) as probe:
            probe.authenticate("user-pw")
            hits=probe[0].search_for("SECRET-123")
            quads=[[*hits[0]]] if hits else None
        self.assertIsNotNone(quads,"样例文档首页应有可检索文字")
        result=self._save_annotations(path,[
            {"action":"add","id":"ul-enc","kind":"underline","page":0,"color":"#FF0000","quads":quads},
        ],password="user-pw")
        self.assertTrue(result["ok"])
        with fitz.open(str(path)) as doc:
            self.assertTrue(doc.needs_pass or doc.is_encrypted)      # 仍加密
            self.assertEqual(doc.authenticate("user-pw"),2)          # 原密码可解
            self.assertEqual(len(doc),3)
            self.assertEqual([a.type[1] for a in doc[0].annots()],["Underline"])
            for index,page in enumerate(doc):
                gray=page.get_pixmap(dpi=36,colorspace=fitz.csGRAY)
                self.assertLess(min(gray.samples),250,f"第 {index+1} 页渲染为空白（内容流被写坏）")

    def test_stamp_annotation_roundtrip(self):
        """V4 T4：预置图章与图片图章——渲染可见性、改色、删除、第三方可见。

        S1 尖刺结论：图片图章若照抄源文档的图片对象字典，其 /ColorSpace 会指向错误对象，
        表现为「注释存在、/AP 存在、**渲染 0 像素**」且无任何报错。因此必须断言渲染真的变了，
        只断言 /Stamp 类型存在会漏掉这种静默失败。
        """
        path=self._annotation_path("annot-stamp.pdf")
        baseline=self._gray_samples(path)
        pix=fitz.Pixmap(fitz.csRGB,fitz.IRect(0,0,40,20))
        pix.set_rect(pix.irect,(200,30,30))
        image_b64=base64.b64encode(pix.tobytes("png")).decode("ascii")

        result=self._save_annotations(path,[
            {"action":"add","id":"st1","kind":"stamp","page":0,"stamp":"approved",
             "color":"#0000FF","rect":[60,80,260,140]},
            {"action":"add","id":"st2","kind":"stamp","page":0,"stamp":"draft",
             "color":"#FF0000","rect":[60,180,260,240]},
            {"action":"add","id":"st3","kind":"stamp","page":0,"image":image_b64,
             "rect":[60,280,260,340]},
        ])
        self.assertTrue(result["ok"])
        for key in ("st1","st2","st3"):
            self.assertIn(key,result["annotationHandles"])

        # 三个图章各自位置都要出现可见变化（图片图章是重点：0 像素即静默失败）
        after=self._gray_samples(path)
        self.assertEqual(self._ranges_changed(baseline,after,[(80,140),(180,240),(280,340)]),
                         [True,True,True],"存在渲染不可见的图章（图片 appearance 流静默失败）")

        with fitz.open(str(path)) as doc:
            self.assertEqual(sorted(a.type[1] for a in doc[0].annots()),["Stamp","Stamp","Stamp"])
            self.assertEqual({a.info.get("name") for a in doc[0].annots()},{"st1","st2","st3"})
            # 图片图章必须是注释（可按名删除），而不是页面内容
            self.assertEqual(len(doc),3)
        reader=PdfReader(str(path))
        subtypes=sorted(a.get_object().get("/Subtype") for pg in reader.pages for a in (pg.get("/Annots") or []))
        self.assertEqual(subtypes,["/Stamp","/Stamp","/Stamp"])

        # 改色（stroke_types 已含 Stamp）+ 删除
        handles=result["annotationHandles"]
        self._save_annotations(path,[{"action":"setColor","name":handles["st1"],"color":"#00FF00"}])
        with fitz.open(str(path)) as doc:
            first=[a for a in doc[0].annots() if a.info.get("name")=="st1"][0]
            self.assertEqual(first.colors.get("stroke"),[0.0,1.0,0.0])
        self._save_annotations(path,[{"action":"delete","name":handles["st2"]}])
        with fitz.open(str(path)) as doc:
            self.assertEqual({a.info.get("name") for a in doc[0].annots()},{"st1","st3"})

    def test_stamp_rejects_bad_token_and_bad_image(self):
        """V4 T4：未知图章 token / 坏图片数据必须显式失败，不得静默产出空注释。"""
        path=self._annotation_path("annot-stamp-bad.pdf")
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"x","kind":"stamp","page":0,
                                          "stamp":"no-such-stamp","rect":[60,80,260,140]}])
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"y","kind":"stamp","page":0,
                                          "image":"bm90LWFuLWltYWdl","rect":[60,80,260,140]}])
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"z","kind":"stamp","page":0,
                                          "stamp":"approved","rect":[60,80,60,140]}])

    def test_stamp_encrypted_save_keeps_content(self):
        """V4 T4：图章在加密文档上保存后仍加密，且逐页内容非空白。"""
        path=self._annotation_path("annot-stamp-enc.pdf",encrypted=True,password="user-pw")
        result=self._save_annotations(path,[
            {"action":"add","id":"st-enc","kind":"stamp","page":0,"stamp":"confidential",
             "rect":[60,80,260,140]},
        ],password="user-pw")
        self.assertTrue(result["ok"])
        with fitz.open(str(path)) as doc:
            self.assertTrue(doc.needs_pass or doc.is_encrypted)
            self.assertEqual(doc.authenticate("user-pw"),2)
            self.assertEqual(len(doc),3)
            self.assertEqual([a.type[1] for a in doc[0].annots()],["Stamp"])
            for index,page in enumerate(doc):
                gray=page.get_pixmap(dpi=36,colorspace=fitz.csGRAY)
                self.assertLess(min(gray.samples),250,f"第 {index+1} 页渲染为空白（内容流被写坏）")

    def test_list_annotations_readback_roundtrip(self):
        """V4 T9a：读回文件内注释——两个名字 / 内容 / 颜色 / 矩形 / 页码，且**只读**。

        关键：读回后源文件字节必须不变（绝不能走 save_annotations 的就地写回路径）。
        """
        path=self._annotation_path("annot-readback.pdf")
        self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"color":"#FF0000","rect":[50,100,300,130]},
            {"action":"add","id":"ul1","kind":"underline","page":1,"color":"#00FF00","quads":[[50,100,300,130]]},
            {"action":"add","id":"n1","kind":"note","page":2,"rect":[50,100,300,130],"text":"读回笔记"},
        ])
        after_save=path.read_bytes()

        result=engine.dispatch({"operation":"list_annotations","files":[str(path)],"options":{}})
        self.assertTrue(result["ok"])
        self.assertEqual(result["pageCount"],3)
        rows=result["annotations"]
        # 笔记 = 关联高亮 + FreeText，故共 4 条
        self.assertEqual(len(rows),4)
        self.assertEqual(sorted(r["subtype"] for r in rows),
                         ["/FreeText","/Highlight","/Highlight","/Underline"])
        self.assertEqual(sorted({r["page"] for r in rows}),[0,1,2])
        # 两个名字都要在：私有 /Name（set_name 真正写入的键）与标准 /NM
        self.assertTrue({"h1","ul1","n1"} <= {r["privateName"] for r in rows},
                        [r["privateName"] for r in rows])
        # 回归锁：稳定名在**私有 /Name**，标准 /NM 是自动名（fitz-A*），二者**不得相等**。
        # 若将来有人把稳定名改写到 /NM（或读错键），这条会立刻失败——
        # 代码注释里也据此写明「按 /NM 匹配会静默错位」。
        for row in rows:
            self.assertNotEqual(row["name"],row["privateName"],
                                f"标准 /NM 不应等于稳定名（row={row['privateName']}）："
                                "稳定名在私有 /Name，/NM 是自动名")
            if row["privateName"] in ("h1","ul1","n1"):
                self.assertRegex(row["name"] or "","^fitz-",
                                 "标准 /NM 应为 PyMuPDF 自动名（fitz-*）")
        for row in rows:
            self.assertIn("name",row)
            self.assertIsNotNone(row["rect"],"每条注释都应带矩形")
        note=[r for r in rows if r["subtype"]=="/FreeText"][0]
        self.assertIn("读回笔记",note["content"] or "")
        highlight=[r for r in rows if r["subtype"]=="/Highlight" and r["page"]==0][0]
        self.assertIn("1 0 0",highlight["color"] or "")

        # **只读**：读回不得改动源文件
        self.assertEqual(path.read_bytes(),after_save)

    def test_list_annotations_pages_filter_and_encrypted(self):
        """V4 T9a：pages 过滤（惰性分页读）+ 加密文档可读回。"""
        path=self._annotation_path("annot-readback-enc.pdf",encrypted=True,password="user-pw")
        self._save_annotations(path,[
            {"action":"add","id":"e1","kind":"highlight","page":0,"rect":[50,100,300,130]},
            {"action":"add","id":"e2","kind":"highlight","page":2,"rect":[50,100,300,130]},
        ],password="user-pw")

        result=engine.dispatch({"operation":"list_annotations","files":[str(path)],
                                "options":{"password":"user-pw","pages":[2]}})
        self.assertTrue(result["ok"])
        rows=result["annotations"]
        self.assertEqual([r["page"] for r in rows],[2])   # 只读第 2 页
        self.assertEqual(rows[0]["privateName"],"e2")

    def test_list_annotations_is_read_only_operation(self):
        """V4 T9a：list_annotations 不进 BY_ID、也不进 ANNOTATION_OPERATIONS（写路径）。"""
        self.assertNotIn("list_annotations",engine.BY_ID)
        self.assertNotIn("list_annotations",engine.ANNOTATION_OPERATIONS)
        self.assertEqual(len(engine.BY_ID),33)

    def test_text_markup_opacity_passthrough(self):
        """V4 T9：属性面板改不透明度 → 引擎透传 set_opacity（文本标记）。"""
        path=self._annotation_path("annot-opacity.pdf")
        self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"rect":[50,100,300,130],"opacity":0.25},
            {"action":"add","id":"u1","kind":"underline","page":1,"quads":[[50,100,300,130]]},
        ])
        with fitz.open(str(path)) as doc:
            values={a.info.get("name"):round(a.opacity,2) for a in doc[0].annots()}
            self.assertEqual(values,{"h1":0.25})                    # 显式覆盖
            # 未给 opacity 的下划线：不调用 set_opacity，PyMuPDF 读回 -1（约定 = 未设置/不透明）
            underline=[a for a in doc[1].annots() if a.type[1]=="Underline"][0]
            self.assertLess(underline.opacity,0)
        # 越界值被钳制到 1.0（不抛错）；PyMuPDF 对「不透明」读回 -1 或 1.0
        self._save_annotations(path,[{"action":"add","id":"h2","kind":"highlight","page":0,
                                      "rect":[50,200,300,230],"opacity":5}])
        with fitz.open(str(path)) as doc:
            over=[a for a in doc[0].annots() if a.info.get("name")=="h2"][0]
            self.assertIn(round(over.opacity,2),(-1.0,1.0))

    def test_reply_annotation_irt_roundtrip(self):
        """V4 T10：回复生成 /IRT 关联注释（标准键；其它阅读器可读），读回通道也带出 irt。"""
        path=self._annotation_path("annot-reply.pdf")
        added=self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"rect":[50,100,300,130]},
        ])
        handle=added["annotationHandles"]["h1"]

        self._save_annotations(path,[
            {"action":"add","id":"r1","kind":"reply","page":0,"parent":handle,"text":"第一条回复"},
            {"action":"add","id":"r2","kind":"reply","page":0,"parent":handle,"text":"第二条回复"},
        ])
        with fitz.open(str(path)) as doc:
            # 注意：必须在迭代内取完所需字段——Annot 对象在 annots() 生成器结束后会失效
            rows={a.info.get("name"):(a.xref,a.type[1],a.info.get("content")) for a in doc[0].annots()}
            self.assertIn("r1",rows);self.assertIn("r2",rows)
            parent_xref=rows["h1"][0]
            for name in ("r1","r2"):
                # 低层 /IRT 必须指向父注释 xref（标准键，非私有）
                self.assertEqual(doc.xref_get_key(rows[name][0],"IRT")[1],f"{parent_xref} 0 R")
            # 回复是 FreeText 且内容正确
            self.assertEqual(rows["r1"][1],"FreeText")
            self.assertIn("第一条回复",rows["r1"][2] or "")

        # 读回通道带出 irt（T9a 的 list_annotations）
        listed=engine.dispatch({"operation":"list_annotations","files":[str(path)],"options":{}})
        replies=[r for r in listed["annotations"] if r["irt"]]
        self.assertEqual(len(replies),2)
        contents={r["privateName"]:r["content"] for r in replies}
        self.assertIn("第一条回复",contents["r1"])

    def test_reply_annotation_validation(self):
        """V4 T10：缺父批注 / 父不存在 / 空内容都必须显式报错。"""
        path=self._annotation_path("annot-reply-bad.pdf")
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"x","kind":"reply","page":0,"text":"无父"}])
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"y","kind":"reply","page":0,
                                          "parent":"no-such-annot","text":"父不存在"}])
        added=self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"rect":[50,100,300,130]}])
        with self.assertRaises(engine.UserError):
            self._save_annotations(path,[{"action":"add","id":"z","kind":"reply","page":0,
                                          "parent":added["annotationHandles"]["h1"],"text":"   "}])

    def test_export_annotations_xfdf_and_summary(self):
        """V4 T11：导出 XFDF + 摘要 PDF；导出源=文件内全部注释；**源文件不变**。"""
        path=self._annotation_path("annot-export.pdf")
        self._save_annotations(path,[
            {"action":"add","id":"h1","kind":"highlight","page":0,"color":"#FF0000","rect":[50,100,300,130]},
            {"action":"add","id":"ul1","kind":"underline","page":1,"color":"#00FF00","quads":[[50,100,300,130]]},
            {"action":"add","id":"n1","kind":"note","page":2,"rect":[50,100,300,130],"text":"导出笔记"},
        ])
        before=path.read_bytes()
        outdir=self.root/"export-out"
        outdir.mkdir(exist_ok=True)

        result=engine.dispatch({"operation":"export_annotations","files":[str(path)],
                                "options":{},"outputDir":str(outdir)})
        self.assertTrue(result["ok"])
        self.assertEqual(result["annotationCount"],4)      # 笔记 = 高亮 + FreeText
        self.assertEqual(len(result["outputs"]),2)         # xfdf + 摘要 PDF

        xfdf=Path([o for o in result["outputs"] if o.endswith(".xfdf")][0])
        summary=Path([o for o in result["outputs"] if o.endswith(".pdf")][0])
        self.assertTrue(xfdf.is_file());self.assertTrue(summary.is_file())

        text=xfdf.read_text(encoding="utf-8")
        self.assertIn("<highlight",text);self.assertIn("<underline",text);self.assertIn("<freetext",text)
        self.assertIn('color="#FF0000"',text)              # /C → #RRGGBB
        self.assertIn('name="h1"',text)                    # 稳定名
        self.assertIn("导出笔记",text)
        self.assertIn('page="2"',text)
        with fitz.open(str(summary)) as doc:
            self.assertGreaterEqual(doc.page_count,1)

        # **只读**：导出不得改动源文件（源文档只读约定）
        self.assertEqual(path.read_bytes(),before)

    def test_export_annotations_includes_external_annotations(self):
        """V4 T11 核心要求：导出源是**文件内全部注释**，不是客户端内存集合——
        其它阅读器（此处用 PyMuPDF 直接写入模拟）标注过的 PDF 也必须导出得出来。"""
        path=self.root/"annot-export-external.pdf"
        doc=fitz.open();page=doc.new_page(width=595,height=842)
        page.insert_text((72,120),"external annotation host",fontsize=12)
        annot=page.add_highlight_annot(quads=[page.search_for("external annotation host")[0]])
        annot.set_info(content="第三方标注",title="Acrobat")
        annot.set_name("EXT-1")
        annot.update()
        doc.save(str(path));doc.close()

        outdir=self.root/"export-out-ext";outdir.mkdir(exist_ok=True)
        result=engine.dispatch({"operation":"export_annotations","files":[str(path)],
                                "options":{"format":"xfdf"},"outputDir":str(outdir)})
        self.assertTrue(result["ok"])
        self.assertEqual(result["annotationCount"],1)
        text=Path(result["outputs"][0]).read_text(encoding="utf-8")
        self.assertIn("<highlight",text)
        self.assertIn("第三方标注",text)
        self.assertIn('title="Acrobat"',text)

    def test_export_annotations_validation(self):
        """V4 T11：多文件 / 缺输出目录 / 非法格式都必须显式报错。"""
        path=self._annotation_path("annot-export-bad.pdf")
        with self.assertRaises(engine.UserError):
            engine.dispatch({"operation":"export_annotations","files":[str(path),str(path)],
                             "options":{},"outputDir":str(self.root)})
        with self.assertRaises(engine.UserError):
            engine.dispatch({"operation":"export_annotations","files":[str(path)],"options":{}})
        with self.assertRaises(engine.UserError):
            engine.dispatch({"operation":"export_annotations","files":[str(path)],
                             "options":{"format":"docx"},"outputDir":str(self.root)})

    def test_export_annotations_stays_out_of_catalog(self):
        """V4 T11：export_annotations 不进 BY_ID（工具目录仍 33）。"""
        self.assertNotIn("export_annotations",engine.BY_ID)
        self.assertEqual(len(engine.BY_ID),33)

    def test_annotation_operations_stay_out_of_catalog(self):
        # 批注 handler 不得进 BY_ID（v1.1 目录为 33，含 translate_batch）
        self.assertEqual(len(engine.BY_ID),33)
        for op in engine.ANNOTATION_OPERATIONS:
            self.assertNotIn(op,engine.BY_ID)
        self.assertIn("add_ink_annot",engine.ANNOTATION_OPERATIONS)
        self.assertIn("add_shape_annot",engine.ANNOTATION_OPERATIONS)
        # 别名 operation 亦可用（不进 catalog）
        alias=self._annotation_path("annot-alias-ink.pdf")
        result=engine.dispatch({"operation":"add_ink_annot","files":[str(alias)],
                                "options":{"id":"a","page":0,"color":"#0000FF",
                                           "strokes":[[[10,10],[60,60]]],"width":2}})
        self.assertTrue(result["ok"])

    def test_edit_tools_retain_ink_and_shape(self):
        source=self._annotation_path("annot-ink-retention.pdf")
        self._save_annotations(source,[
            {"action":"add","id":"i","kind":"ink","page":0,"color":"#0000FF",
             "strokes":[[[60,120],[160,180]]],"width":3},
            {"action":"add","id":"s","kind":"shape","shape":"square","page":0,"color":"#00FF00",
             "rect":[60,300,220,360],"width":2},
            {"action":"add","id":"ci","kind":"shape","shape":"circle","page":1,"color":"#FF00FF",
             "rect":[60,300,220,360],"width":2},
        ])
        for operation,options in [("compress",{"quality":"balanced"}),
                                  ("rotate",{"angle":"90","pages":"1"})]:
            result=self.run_tool(operation,options,files=[source])
            types=sorted(self._annot_types(result["outputs"][0]))
            self.assertEqual(types,["Circle","Ink","Square"],f"{operation} 改变注释类型/数量：{types}")

    # --- V2 T10 FR-23：编辑类工具注释保留（Q2 修订：compress/crop/rotate 保留，redact 删除全部） ---

    def _annotated_retention_pdf(self,name="annot-retention.pdf"):
        """合成含 2 Highlight + 1 FreeText 的 PDF；另含 1 个 Widget 表单字段与 TOC（供 redact M-12 断言）。

        FreeText 放置在页面下部（y≈700），裁剪到 600pt 高后不再可见，用于验证 crop 不物理删除区域外注释（Q8）。
        """
        path=self.root/name
        doc=fitz.open(self.source)  # 3 页，含 SECRET-123 文本（redact 命中）
        widget=fitz.Widget()
        widget.field_name="aipdf_retention_field"
        widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT
        widget.rect=fitz.Rect(50,700,250,730)
        widget.field_value="Ada Lovelace"
        doc[0].add_widget(widget)
        doc.set_toc([[1,"Chapter 1",1],[1,"Chapter 2",2]])
        engine.add_highlight_annot(doc[0],rect=fitz.Rect(50,80,300,110))
        engine.add_highlight_annot(doc[1],rect=fitz.Rect(50,80,300,110))
        engine.add_freetext_annot(doc[2],rect=fitz.Rect(50,700,300,740),text="保留注释")
        engine.save_pdf(doc,path)
        doc.close()
        return path

    @staticmethod
    def _annot_types(path):
        with fitz.open(str(path)) as doc:
            return [a.type[1] for page in doc for a in (page.annots() or [])]

    def test_edit_tools_retain_annotations(self):
        source=self._annotated_retention_pdf()
        original=self._annot_types(source)
        self.assertEqual(sorted(original),["FreeText","Highlight","Highlight"])
        for operation,options in [("compress",{"quality":"balanced"}),
                                  ("crop",{"x":"25","y":"30","width":"400","height":"600"}),
                                  ("rotate",{"angle":"90","pages":"1"})]:
            result=self.run_tool(operation,options,files=[source])
            types=self._annot_types(result["outputs"][0])
            self.assertEqual(len(types),3,f"{operation} 改变注释数量：{types}")
            self.assertEqual(types.count("Highlight"),2,f"{operation} 改变 Highlight 数量")
            self.assertEqual(types.count("FreeText"),1,f"{operation} 改变 FreeText 数量")

    def test_crop_keeps_out_of_box_annotations(self):
        source=self._annotated_retention_pdf()
        # FreeText 位于 y≈700，裁剪到 600pt 高后不再可见，但必须仍在文件中（Q8 不物理删除）
        result=self.run_tool("crop",{"x":"25","y":"30","width":"400","height":"600"},files=[source])
        types=self._annot_types(result["outputs"][0])
        self.assertIn("FreeText",types)
        self.assertEqual(len(types),3)

    def test_redact_removes_annotations_widgets_and_toc(self):
        source=self._annotated_retention_pdf()
        with fitz.open(str(source)) as doc:
            self.assertEqual(len(doc.get_toc()),2)
            self.assertTrue([w for p in doc for w in (p.widgets() or [])])
        result=self.run_tool("redact",{"text":"SECRET-123"},files=[source])
        with fitz.open(str(result["outputs"][0])) as doc:
            self.assertEqual([a for p in doc for a in (p.annots() or [])],[])   # 注释归零（Q2）
            self.assertEqual([w for p in doc for w in (p.widgets() or [])],[])  # Widget 清除（M-12）
            self.assertEqual(doc.get_toc(),[])                                   # TOC 清空（M-12）

    def test_corrupt_input_does_not_harm_original(self):
        broken=self.root/"broken-annot.pdf";broken.write_bytes(b"%PDF-1.7\nnot really a pdf")
        before=hashlib.sha256(broken.read_bytes()).hexdigest()
        with self.assertRaises(engine.UserError):self._save_annotations(broken,[{"action":"add","kind":"highlight","page":0,"rect":[0,0,10,10]}])
        self.assertEqual(hashlib.sha256(broken.read_bytes()).hexdigest(),before)
        truncated=self.root/"truncated-annot.pdf"
        data=self.source.read_bytes();truncated.write_bytes(data[:len(data)//2])
        before2=hashlib.sha256(truncated.read_bytes()).hexdigest()
        with self.assertRaises(engine.UserError):self._save_annotations(truncated,[{"action":"add","kind":"highlight","page":0,"rect":[0,0,10,10]}])
        self.assertEqual(hashlib.sha256(truncated.read_bytes()).hexdigest(),before2)

    def test_encrypt_unlock(self):
        protected=self.run_tool("protect",{"newPassword":"local-test-password"})
        with fitz.open(protected["outputs"][0]) as doc:self.assertTrue(doc.needs_pass)
        with self.assertRaises(engine.UserError):self.run_tool("unlock",files=protected["outputs"])
        with self.pdf(self.run_tool("unlock",{"password":"local-test-password"},protected["outputs"])) as doc:
            self.assertFalse(doc.needs_pass);self.assertIn("Page 1",doc[0].get_text())

    def test_compression_retains_text(self):
        for quality in ["lossless","balanced","small"]:
            with self.pdf(self.run_tool("compress",{"quality":quality})) as d:self.assertIn("Unique content",d[0].get_text())
        image_pdf=self.run_tool("images_to_pdf",files=[self.image])
        result=self.run_tool("compress",{"quality":"small"},image_pdf["outputs"])
        self.assertLess(result["outputBytes"],result["inputBytes"])

    def test_repair(self):
        damaged=self.root/"damaged.pdf"
        data=self.source.read_bytes()
        damaged.write_bytes(data[:data.rfind(b"startxref")]+b"startxref\n0\n%%EOF\n")
        with self.pdf(self.run_tool("repair",files=[damaged])) as d:self.assertEqual(len(d),3)

    def test_images_roundtrip(self):
        result=self.run_tool("pdf_to_images",{"format":"png","pages":"1-2"})
        self.assertEqual(len(result["outputs"]),2)
        result=self.run_tool("images_to_pdf",files=result["outputs"])
        with self.pdf(result) as d:self.assertEqual(len(d),2)
        embedded=self.run_tool("pdf_to_images",{"mode":"embedded"},result["outputs"])
        self.assertEqual(len(embedded["outputs"]),2)

    def test_office_exports(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        docx=self.run_tool("pdf_to_word")
        document=Document(docx["outputs"][0]);self.assertTrue(document.tables)
        self.assertIn("Page 1",' '.join(p.text for p in document.paragraphs))
        visual=self.run_tool("pdf_to_word",{"mode":"visual"})
        self.assertEqual(len(Document(visual["outputs"][0]).inline_shapes),3)
        for mode in ["visual","editable"]:
            ppt=self.run_tool("pdf_to_ppt",{"mode":mode})
            self.assertEqual(len(Presentation(ppt["outputs"][0]).slides),3)
        xlsx=self.run_tool("pdf_to_excel")
        workbook=load_workbook(xlsx["outputs"][0]);self.assertEqual(len(workbook.sheetnames),3)
        self.assertEqual(workbook.active.cell(2,2).value,"100")

    def test_office_imports_and_pdfa(self):
        if not engine.find_office():self.skipTest("LibreOffice not installed")
        for out,incoming in [("pdf_to_word","word_to_pdf"),("pdf_to_ppt","ppt_to_pdf"),("pdf_to_excel","excel_to_pdf")]:
            source=self.run_tool(out)
            with self.pdf(self.run_tool(incoming,files=source["outputs"])) as d:self.assertGreater(len(d),0)
        with self.pdf(self.run_tool("pdfa")) as d:self.assertIn("pdfaid",d.get_xml_metadata())

    def test_markdown(self):
        path=self.run_tool("markdown")["outputs"][0]
        text=Path(path).read_text(encoding="utf-8")
        self.assertIn("# AIPDF",text);self.assertIn("| Item | Value |",text);self.assertIn("本地",text)

    def test_html(self):
        html=self.root/"test.html";html.write_text("<html><body><h1>Local HTML</h1><p>中文网页</p></body></html>",encoding="utf-8")
        with self.pdf(self.run_tool("html_to_pdf",files=[html])) as d:
            self.assertIn("Local HTML",d[0].get_text());self.assertIn("中文网页",d[0].get_text())

    def test_compare(self):
        result=self.run_tool("compare",files=[self.source,self.second])
        self.assertEqual(len(result["outputs"]),2)
        self.assertIn("Changed",Path(result["outputs"][1]).read_text())
        with self.assertRaises(engine.UserError):self.run_tool("compare")

    def test_create_and_fill_forms_canonical_tree(self):
        result=self.run_tool("forms",{"formMode":"create","fieldName":"person","fieldValue":"Original","x":"50","y":"400","width":"300"})
        path=result["outputs"][0]
        self.assertIn("person",PdfReader(path).get_fields())
        filled=self.run_tool("forms",{"formMode":"fill","formValues":{"person":"Ada Lovelace"}},[path])
        reader=PdfReader(filled["outputs"][0])
        self.assertEqual(reader.get_fields()["person"]["/V"],"Ada Lovelace")
        widget=reader.pages[0]["/Annots"][0].get_object()
        self.assertEqual(widget["/V"],"Ada Lovelace")
        self.assertTrue(widget["/AP"]["/N"])
        with self.pdf(filled) as d:self.assertIn("Ada Lovelace",d[0].get_text())
        for kind in ["checkbox","list","radio"]:
            result=self.run_tool("forms",{"formMode":"create","fieldType":kind,"fieldName":kind,"choices":"One\nTwo","fieldValue":"One"})
            self.assertIn(kind,PdfReader(result["outputs"][0]).get_fields())

    def test_workflow(self):
        result=self.run_tool("workflow",{"steps":[{"operation":"rotate","options":{"angle":"90"}},{"operation":"extract","options":{"pages":"2"}},{"operation":"protect","options":{"newPassword":"abc"}},{"operation":"unlock","options":{}}]})
        with self.pdf(result) as d:self.assertEqual(len(d),1);self.assertEqual(d[0].rotation,90);self.assertFalse(d.needs_pass)

    def test_removed_ai_tools_and_offline_health(self):
        self.assertEqual(len(engine.BY_ID),33)
        self.assertEqual(engine.BY_ID["markdown"]["category"],"格式转换")
        # v1.1：翻译正式纳入范围。目录 id 为 translate_batch，请求 operation 为
        # translate，专用 dispatch 处理，二者都不应被当成「未知的工具」。
        self.assertEqual(engine.BY_ID["translate_batch"]["name"],"批量翻译")
        self.assertNotIn("translate",engine.BY_ID)
        with self.assertRaisesRegex(engine.UserError,"翻译服务端点"):
            self.run_tool("translate")
        # AI 摘要仍由维护者排除。
        self.assertNotIn("summarize",engine.BY_ID)
        with self.assertRaisesRegex(engine.UserError,"未知的工具"):
            self.run_tool("summarize")
        with patch("socket.create_connection",side_effect=AssertionError("Health must not connect to services")):
            status=engine.health()
        self.assertTrue(status["ok"])
        self.assertNotIn("models",status)
        self.assertNotIn("ollama",status)

    def test_worker_json_protocol(self):
        for op in ["rotate","pdf_to_word","pdf_to_excel","markdown"]:
            request={"operation":op,"files":[str(self.source)],"options":{"angle":"90"},"outputDir":str(self.root)}
            result=subprocess.run([sys.executable,str(ROOT/"backend/engine.py")],input=json.dumps(request),capture_output=True,text=True)
            self.assertEqual(result.returncode,0)
            self.assertTrue(json.loads(result.stdout)["ok"])
            self.assertIn("message",json.loads(result.stderr.splitlines()[0]))

    def test_cjk_forms_and_merge_field_names(self):
        form=self.run_tool("forms",{"formMode":"create","fieldName":"person","fieldValue":"张三"})
        with self.pdf(form) as d:self.assertIn("张三",d[0].get_text())
        merged=self.run_tool("merge",files=form["outputs"]*2)
        fields=PdfReader(merged["outputs"][0]).get_fields()
        self.assertEqual(len(fields),2)
        filled=self.run_tool("forms",{"formMode":"fill","formValues":{"person":"李四"}},form["outputs"])
        self.assertEqual(PdfReader(filled["outputs"][0]).get_fields()["person"]["/V"],"李四")

    def test_web_url(self):
        if not engine.web_path():self.skipTest("Web helper not built")
        import http.server
        import threading
        class Handler(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(200);self.send_header("Content-Type","text/html; charset=utf-8");self.end_headers()
                self.wfile.write(b"<html><body><h1>Local Web Test</h1><p>Browser rendered PDF.</p></body></html>")
            def log_message(self,*args):pass
        server=http.server.HTTPServer(("127.0.0.1",0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            result=engine.dispatch({"operation":"html_to_pdf","files":[],"options":{"url":f"http://127.0.0.1:{server.server_port}/"},"outputDir":str(self.root)})
            with self.pdf(result) as d:self.assertIn("Local Web Test",d[0].get_text())
        finally:
            server.shutdown();server.server_close();thread.join()

    def test_cancel_stops_children(self):
        if sys.platform != "darwin":
            self.skipTest("Mac-only test: uses /usr/bin/pgrep")
        import signal
        import time
        helper=self.root/"slow-helper"
        helper.write_text("#!/bin/sh\nexec /bin/sleep 120\n")
        helper.chmod(0o755)
        request={"operation":"scan","files":[str(self.image)],"options":{},"outputDir":str(self.root)}
        env={**os.environ,"AIPDF_VISION":str(helper)}
        p=subprocess.Popen([sys.executable,str(ROOT/"backend/engine.py")],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,env=env)
        p.stdin.write(json.dumps(request));p.stdin.close();p.stdin=None
        # Wait until the child exists, then exercise the same SIGTERM used by the UI.
        child_pid=None
        for _ in range(100):
            info=subprocess.run(["/usr/bin/pgrep","-P",str(p.pid)],capture_output=True,text=True)
            if info.stdout.strip():child_pid=int(info.stdout.splitlines()[0]);break
            time.sleep(.02)
        self.assertIsNotNone(child_pid)
        p.terminate();out,_=p.communicate(timeout=5)
        self.assertFalse(json.loads(out)["ok"])
        with self.assertRaises(ProcessLookupError):os.kill(child_pid,0)

    def test_vision_ocr_scan(self):
        if not engine.vision_path():self.skipTest("Vision helper not built")
        scan=self.run_tool("scan",files=[self.image])
        with self.pdf(scan) as d:self.assertEqual(len(d),1)
        result=self.run_tool("ocr",files=scan["outputs"])
        with self.pdf(result) as d:
            text=d[0].get_text()
            self.assertIn("AIPDF",text)
            self.assertIn("PDF",text)

    def test_bounds_and_invalid_inputs(self):
        for options in [{"width":"0"},{"x":"-1"},{"x":"inf"},{"height":"99999"}]:
            with self.assertRaises(engine.UserError):self.run_tool("crop",options)
        with self.assertRaises(engine.UserError):self.run_tool("images_to_pdf")
        with self.assertRaises(engine.UserError):self.run_tool("edit",{"text":"A"*10000})


if __name__=="__main__":
    (ROOT/"tmp").mkdir(exist_ok=True)
    unittest.main(verbosity=2)
